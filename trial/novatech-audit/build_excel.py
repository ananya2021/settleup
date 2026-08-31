#!/usr/bin/env python3
"""Build the NovaTech Industries audit Excel workbook."""

import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

# ── Style definitions ──
NAVY = "1F3864"
TEAL = "0D4F4F"
WHITE = "FFFFFF"
LIGHT_GRAY = "F2F2F2"
DARK_GRAY = "404040"

title_font = Font(name="Calibri", size=16, bold=True, color=WHITE)
header_font = Font(name="Calibri", size=11, bold=True, color=WHITE)
data_font = Font(name="Calibri", size=10, color=DARK_GRAY)
bold_font = Font(name="Calibri", size=10, bold=True, color=DARK_GRAY)
section_font = Font(name="Calibri", size=12, bold=True, color=NAVY)

title_fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
header_fill = PatternFill(start_color=TEAL, end_color=TEAL, fill_type="solid")
alt_fill = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")
critical_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
high_fill = PatternFill(start_color="FF8C00", end_color="FF8C00", fill_type="solid")
medium_fill = PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid")
low_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")

thin_border = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
wrap_align = Alignment(wrap_text=True, vertical="top")


def load(fn):
    with open(fn) as f:
        return json.load(f)


def write_headers(ws, row, headers, col_start=1):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=col_start + i, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center_align
        c.border = thin_border


def write_row(ws, row, values, col_start=1):
    for i, v in enumerate(values):
        if isinstance(v, (list, dict)):
            v = str(v)
        c = ws.cell(row=row, column=col_start + i, value=v)
        c.font = data_font
        c.border = thin_border
        c.alignment = left_align
        if row % 2 == 0:
            c.fill = alt_fill


def write_title(ws, title, ncols=8):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=title)
    c.font = title_font
    c.fill = title_fill
    c.alignment = center_align
    ws.row_dimensions[1].height = 36


def auto_width(ws, min_w=10, max_w=38):
    for col_cells in ws.columns:
        mx = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            try:
                mx = max(mx, min(len(str(cell.value or "")), max_w))
            except:
                pass
        ws.column_dimensions[col_letter].width = max(mx + 2, min_w)


def freeze_top(ws, row=3):
    ws.freeze_panes = ws.cell(row=row, column=1)


# ═══════════════════════════════════════════════════════════
# Load all data
# ═══════════════════════════════════════════════════════════
B = "."
employees = load(f"{B}/master-data/employees.json")
vendors = load(f"{B}/master-data/vendors.json")
customers = load(f"{B}/master-data/customers.json")
products = load(f"{B}/master-data/products.json")
warehouses = load(f"{B}/master-data/warehouses.json")
departments = load(f"{B}/master-data/departments.json")
po = load(f"{B}/transaction-data/purchase-orders.json")
pi = load(f"{B}/transaction-data/purchase-invoices.json")
so = load(f"{B}/transaction-data/sales-orders.json")
si = load(f"{B}/transaction-data/sales-invoices.json")
expenses = load(f"{B}/transaction-data/expense-claims.json")
bank = load(f"{B}/transaction-data/bank-transactions.json")
inv = load(f"{B}/transaction-data/inventory-movements.json")
payroll = load(f"{B}/transaction-data/payroll-records.json")
tax = load(f"{B}/transaction-data/tax-payments.json")
budgets = load(f"{B}/transaction-data/budget-vs-actuals.json")
uaccess = load(f"{B}/compliance-data/user-access-records.json")
login = load(f"{B}/compliance-data/login-activity.json")
awf = load(f"{B}/compliance-data/approval-workflows.json")
findings_data = load(f"{B}/audit-findings/findings.json")
anomalies_data = load(f"{B}/audit-findings/anomaly-detection.json")

findings = findings_data["findings"]
anomaly_cats = anomalies_data.get("anomaly_categories", [])
susp_entities = anomalies_data.get("suspicious_entities", [])
fraud_indicators = anomalies_data.get("fraud_indicators", [])

print("All data loaded.")


# ═══════════════════════════════════════════════════════════
# SHEET 1: Dashboard
# ═══════════════════════════════════════════════════════════
ws = wb.active
ws.title = "Dashboard"
ws.sheet_properties.tabColor = NAVY

ws.merge_cells("A1:J1")
c = ws.cell(row=1, column=1, value="NovaTech Industries Pvt. Ltd. — Audit Dashboard")
c.font = Font(name="Calibri", size=20, bold=True, color=WHITE)
c.fill = title_fill
c.alignment = center_align
ws.row_dimensions[1].height = 50

ws.merge_cells("A2:J2")
c = ws.cell(row=2, column=1, value="Comprehensive Internal Audit  |  FY 2025-26  |  Overall Rating: C — Needs Significant Improvement")
c.font = Font(name="Calibri", size=11, italic=True, color=TEAL)
c.alignment = center_align
ws.row_dimensions[2].height = 28

# Dataset overview
r = 4
ws.cell(row=r, column=1, value="DATASET OVERVIEW").font = section_font
ws.merge_cells(f"A{r}:D{r}")
r += 1
write_headers(ws, r, ["Category", "Count", "Category", "Count"])
r += 1
rows_data = [
    ["Employees", len(employees), "Vendors", len(vendors)],
    ["Customers", len(customers), "Products", len(products)],
    ["Purchase Orders", len(po), "Purchase Invoices", len(pi)],
    ["Sales Orders", len(so), "Sales Invoices", len(si)],
    ["Expense Claims", len(expenses), "Bank Transactions", len(bank)],
    ["Inventory Movements", len(inv), "Payroll Records", len(payroll)],
    ["Tax Payments", len(tax), "User Access Records", len(uaccess)],
    ["Login Sessions", len(login), "Departments", len(departments)],
    ["Warehouses", len(warehouses), "Budget Items", len(budgets)],
]
for rd in rows_data:
    write_row(ws, r, rd)
    r += 1

# Severity breakdown
r += 1
ws.cell(row=r, column=1, value="FINDINGS BY SEVERITY").font = section_font
ws.merge_cells(f"A{r}:D{r}")
r += 1
write_headers(ws, r, ["Severity", "Count", "Financial Impact (₹)", ""])
r += 1
sev_map = {}
for f in findings:
    s = f.get("severity", "Unknown")
    sev_map[s] = sev_map.get(s, 0) + 1
total_impact = 0
for s in ["Critical", "High", "Medium", "Low", "Observation"]:
    cnt = sev_map.get(s, 0)
    if cnt == 0:
        continue
    imp = sum(f.get("financial_impact", 0) for f in findings if f.get("severity") == s)
    total_impact += imp
    write_row(ws, r, [s, cnt, f"₹{imp:,.0f}"])
    cell = ws.cell(row=r, column=1)
    if s == "Critical":
        cell.fill = critical_fill
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    elif s == "High":
        cell.fill = high_fill
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    elif s == "Medium":
        cell.fill = medium_fill
        cell.font = Font(name="Calibri", size=10, bold=True)
    elif s == "Low":
        cell.fill = low_fill
    r += 1
write_row(ws, r, ["TOTAL", len(findings), f"₹{total_impact:,.0f}"])
ws.cell(row=r, column=1).font = bold_font
ws.cell(row=r, column=2).font = bold_font

# Audit area breakdown
r += 2
ws.cell(row=r, column=1, value="FINDINGS BY AUDIT AREA").font = section_font
ws.merge_cells(f"A{r}:D{r}")
r += 1
write_headers(ws, r, ["Audit Area", "Findings", "Critical/High", ""])
r += 1
area_map = {}
for f in findings:
    a = f.get("area", "Unknown")
    area_map[a] = area_map.get(a, 0) + 1
for area, cnt in sorted(area_map.items(), key=lambda x: -x[1]):
    ch = sum(1 for f in findings if f.get("area") == area and f.get("severity") in ("Critical", "High"))
    write_row(ws, r, [area, cnt, ch])
    r += 1

ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 24
ws.column_dimensions["D"].width = 16
print("  Dashboard ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Employees
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Employees")
ws.sheet_properties.tabColor = "4472C4"
write_title(ws, "Employee Master Data", 12)
headers = ["emp_id", "first_name", "last_name", "department", "designation", "grade",
           "monthly_ctc", "date_of_joining", "status", "location", "reporting_to", "pf_number"]
write_headers(ws, 3, headers)
for i, e in enumerate(employees):
    r = 4 + i
    write_row(ws, r, [e.get(h, "") for h in headers])
auto_width(ws)
freeze_top(ws, 4)
print(f"  Employees ({len(employees)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Vendors
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Vendors")
ws.sheet_properties.tabColor = "ED7D31"
write_title(ws, "Vendor Master Data", 12)
headers = ["vendor_id", "vendor_name", "category", "city", "state", "gst_number",
           "pan", "status", "payment_terms", "credit_limit", "ytd_purchases", "last_transaction_date"]
write_headers(ws, 3, headers)
for i, v in enumerate(vendors):
    write_row(ws, 4 + i, [v.get(h, "") for h in headers])
auto_width(ws)
freeze_top(ws, 4)
print(f"  Vendors ({len(vendors)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Customers
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Customers")
ws.sheet_properties.tabColor = "70AD47"
write_title(ws, "Customer Master Data", 12)
headers = ["customer_id", "customer_name", "segment", "city", "state", "region",
           "gst_number", "credit_limit", "outstanding_balance", "ytd_revenue", "status", "account_manager"]
write_headers(ws, 3, headers)
for i, c in enumerate(customers):
    write_row(ws, 4 + i, [c.get(h, "") for h in headers])
auto_width(ws)
freeze_top(ws, 4)
print(f"  Customers ({len(customers)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Products
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Products")
ws.sheet_properties.tabColor = "FFC000"
write_title(ws, "Product Master", 10)
headers = ["sku", "product_name", "category", "unit", "cost_price", "selling_price",
           "min_stock", "max_stock", "lead_time_days", "status"]
write_headers(ws, 3, headers)
for i, p in enumerate(products):
    write_row(ws, 4 + i, [p.get(h, "") for h in headers])
auto_width(ws)
freeze_top(ws, 4)
print(f"  Products ({len(products)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Warehouses
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Warehouses")
ws.sheet_properties.tabColor = "5B9BD5"
write_title(ws, "Warehouse Master", 8)
headers = ["warehouse_id", "name", "city", "state", "capacity_sqft", "manager_id", "total_stock_value", "status"]
write_headers(ws, 3, headers)
for i, w in enumerate(warehouses):
    write_row(ws, 4 + i, [w.get(h, "") for h in headers])
auto_width(ws)
freeze_top(ws, 4)
print(f"  Warehouses ({len(warehouses)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Departments
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Departments")
ws.sheet_properties.tabColor = "A5A5A5"
write_title(ws, "Department Master", 6)
headers = ["dept_code", "department", "budget_code", "annual_budget", "location", "head_count"]
write_headers(ws, 3, headers)
for i, d in enumerate(departments):
    write_row(ws, 4 + i, [d.get(h, "") for h in headers])
auto_width(ws)
freeze_top(ws, 4)
print(f"  Departments ({len(departments)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Purchase Orders
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Purchase Orders")
ws.sheet_properties.tabColor = "ED7D31"
write_title(ws, "Purchase Orders", 14)
headers = ["po_id", "po_date", "vendor_id", "vendor_name", "item_description", "quantity",
           "unit_price", "total_amount", "expected_delivery", "status", "approved_by",
           "approval_date", "warehouse", "plant"]
write_headers(ws, 3, headers)
for i, p in enumerate(po):
    write_row(ws, 4 + i, [p.get(h, "") for h in headers])
auto_width(ws)
freeze_top(ws, 4)
print(f"  Purchase Orders ({len(po)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Purchase Invoices
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Purchase Invoices")
ws.sheet_properties.tabColor = "C55A11"
write_title(ws, "Purchase Invoices", 14)
headers = ["pi_id", "pi_date", "vendor_id", "vendor_name", "po_reference", "item_description",
           "quantity", "unit_price", "total_amount", "gst_amount", "net_amount",
           "payment_status", "payment_date", "supporting_docs"]
write_headers(ws, 3, headers)
for i, p in enumerate(pi):
    write_row(ws, 4 + i, [p.get(h, "") for h in headers])
auto_width(ws)
freeze_top(ws, 4)
print(f"  Purchase Invoices ({len(pi)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Sales Orders
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Sales Orders")
ws.sheet_properties.tabColor = "70AD47"
write_title(ws, "Sales Orders", 16)
headers = ["so_id", "so_date", "customer_id", "customer_name", "item_description", "sku",
           "quantity", "unit_price", "discount_pct", "total_amount", "gst_amount",
           "net_amount", "delivery_date", "status", "sales_rep", "warehouse"]
write_headers(ws, 3, headers)
for i, s in enumerate(so):
    write_row(ws, 4 + i, [s.get(h, "") for h in headers])
auto_width(ws)
freeze_top(ws, 4)
print(f"  Sales Orders ({len(so)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Sales Invoices
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Sales Invoices")
ws.sheet_properties.tabColor = "548235"
write_title(ws, "Sales Invoices", 16)
headers = ["si_id", "si_date", "customer_id", "customer_name", "so_reference", "item_description",
           "sku", "quantity", "unit_price", "discount_pct", "total_amount", "gst_amount",
           "net_amount", "payment_status", "due_date", "payment_received_date"]
write_headers(ws, 3, headers)
for i, s in enumerate(si):
    write_row(ws, 4 + i, [s.get(h, "") for h in headers])
auto_width(ws)
freeze_top(ws, 4)
print(f"  Sales Invoices ({len(si)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Expense Claims
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Expense Claims")
ws.sheet_properties.tabColor = "BF8F00"
write_title(ws, "Expense Claims", 14)
headers = ["claim_id", "claim_date", "employee_id", "employee_name", "department", "category",
           "description", "amount", "claimed_amount", "approved_amount", "approved_by",
           "status", "supporting_docs", "payment_status"]
write_headers(ws, 3, headers)
for i, e in enumerate(expenses):
    write_row(ws, 4 + i, [e.get(h, "") for h in headers])
auto_width(ws)
freeze_top(ws, 4)
print(f"  Expense Claims ({len(expenses)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Bank Transactions
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Bank Transactions")
ws.sheet_properties.tabColor = "44546A"
write_title(ws, "Bank Transactions", 12)
headers = ["txn_id", "txn_date", "txn_time", "type", "description", "reference",
           "debit", "credit", "balance", "bank_account", "reconciled", "reconciled_date"]
write_headers(ws, 3, headers)
for i, b in enumerate(bank):
    write_row(ws, 4 + i, [b.get(h, "") for h in headers])
auto_width(ws)
freeze_top(ws, 4)
print(f"  Bank Transactions ({len(bank)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Inventory Movements
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Inventory Movements")
ws.sheet_properties.tabColor = "9DC3E6"
write_title(ws, "Inventory Movements", 14)
headers = ["movement_id", "movement_date", "movement_type", "sku", "product_name",
           "warehouse_id", "warehouse_name", "quantity", "unit_cost", "total_value",
           "reference_id", "reference_type", "authorized_by", "remarks"]
write_headers(ws, 3, headers)
for i, m in enumerate(inv):
    write_row(ws, 4 + i, [m.get(h, "") for h in headers])
auto_width(ws)
freeze_top(ws, 4)
print(f"  Inventory Movements ({len(inv)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Payroll
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Payroll")
ws.sheet_properties.tabColor = "7030A0"
write_title(ws, "Payroll Records", 16)
headers = ["payroll_id", "employee_id", "employee_name", "department", "month", "year",
           "basic_salary", "hra", "conveyance", "medical", "special_allowance",
           "gross_salary", "pf_deduction", "esi_deduction", "tds_deduction", "net_salary"]
write_headers(ws, 3, headers)
for i, p in enumerate(payroll):
    write_row(ws, 4 + i, [p.get(h, "") for h in headers])
auto_width(ws)
freeze_top(ws, 4)
print(f"  Payroll ({len(payroll)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Tax Payments
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Tax Payments")
ws.sheet_properties.tabColor = "FF0000"
write_title(ws, "Tax Payments", 10)
headers = ["tax_id", "tax_type", "period", "due_date", "payment_date", "amount",
           "penalty", "total_paid", "reference_number", "status"]
write_headers(ws, 3, headers)
for i, t in enumerate(tax):
    write_row(ws, 4 + i, [t.get(h, "") for h in headers])
auto_width(ws)
freeze_top(ws, 4)
print(f"  Tax Payments ({len(tax)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Budget vs Actuals
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Budget vs Actuals")
ws.sheet_properties.tabColor = "00B050"
write_title(ws, "Budget vs Actuals — FY 2025-26", 10)
headers = ["department", "budget_code", "annual_budget", "actual_spend", "variance",
           "variance_percent", "status", "q1_budget", "q2_budget", "q3_budget"]
write_headers(ws, 3, headers)
for i, b in enumerate(budgets):
    write_row(ws, 4 + i, [b.get(h, "") for h in headers])
auto_width(ws)
freeze_top(ws, 4)
print(f"  Budget vs Actuals ({len(budgets)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: User Access
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("User Access")
ws.sheet_properties.tabColor = "FF6600"
write_title(ws, "IT & User Access Records", 14)
headers = ["record_id", "user_id", "employee_id", "employee_name", "department",
           "user_role", "system", "access_level", "is_admin", "mfa_enabled",
           "last_login", "status", "created_date", "created_by"]
write_headers(ws, 3, headers)
for i, u in enumerate(uaccess):
    write_row(ws, 4 + i, [u.get(h, "") for h in headers])
auto_width(ws)
freeze_top(ws, 4)
print(f"  User Access ({len(uaccess)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Login Activity
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Login Activity")
ws.sheet_properties.tabColor = "808000"
write_title(ws, "Login Activity Log", 12)
headers = ["log_id", "user_id", "employee_name", "system", "login_date", "login_time",
           "logout_time", "ip_address", "location", "device", "status", "remarks"]
write_headers(ws, 3, headers)
for i, l in enumerate(login):
    write_row(ws, 4 + i, [l.get(h, "") for h in headers])
auto_width(ws)
freeze_top(ws, 4)
print(f"  Login Activity ({len(login)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Approval Workflows
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Approval Workflows")
ws.sheet_properties.tabColor = "0070C0"
write_title(ws, "Internal Controls — Approval Workflows", 8)

# Approval limits
r = 3
ws.cell(row=r, column=1, value="APPROVAL LIMITS").font = section_font
ws.merge_cells(f"A{r}:D{r}")
r += 1
write_headers(ws, r, ["Role", "Approval Limit (₹)", "Areas", ""])
r += 1
for al in awf.get("approval_limits", []):
    areas_str = ", ".join(al.get("areas", []))
    write_row(ws, r, [al.get("role", ""), al.get("approval_limit", ""), areas_str])
    r += 1

# Segregation of duties
r += 1
ws.cell(row=r, column=1, value="SEGREGATION OF DUTIES RULES").font = section_font
ws.merge_cells(f"A{r}:D{r}")
r += 1
write_headers(ws, r, ["Rule", "Description", "Violations", ""])
r += 1
for sod in awf.get("segregation_of_duties", []):
    write_row(ws, r, [sod.get("rule", ""), sod.get("description", ""), sod.get("violations", 0)])
    r += 1

# Control tests
r += 1
ws.cell(row=r, column=1, value="CONTROL TESTS").font = section_font
ws.merge_cells(f"A{r}:D{r}")
r += 1
write_headers(ws, r, ["Control Area", "Test", "Result", "Details"])
r += 1
for ct in awf.get("control_tests", []):
    write_row(ws, r, [ct.get("area", ""), ct.get("test", ""), ct.get("result", ""), ct.get("details", "")])
    r += 1

auto_width(ws)
freeze_top(ws, 5)
print("  Approval Workflows ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Audit Findings
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Audit Findings")
ws.sheet_properties.tabColor = "FF0000"
write_title(ws, "Audit Findings Register", 18)
headers = ["finding_id", "area", "title", "severity", "risk_score", "financial_impact",
           "reference", "description", "expected_control", "observation", "risk",
           "root_cause", "recommendation", "management_response", "responsible_dept",
           "responsible_person", "due_date", "status"]
write_headers(ws, 3, headers)
for i, f in enumerate(findings):
    r = 4 + i
    vals = [f.get(h, "") for h in headers]
    write_row(ws, r, vals)
    # Color severity
    sev_cell = ws.cell(row=r, column=4)
    sev = f.get("severity", "")
    if sev == "Critical":
        sev_cell.fill = critical_fill
        sev_cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    elif sev == "High":
        sev_cell.fill = high_fill
        sev_cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    elif sev == "Medium":
        sev_cell.fill = medium_fill
        sev_cell.font = Font(name="Calibri", size=10, bold=True)
    elif sev == "Low":
        sev_cell.fill = low_fill

auto_width(ws)
# Make wide columns wider for readability
for col_letter in ["C", "I", "J", "L", "M"]:
    ws.column_dimensions[col_letter].width = 50
ws.column_dimensions["H"].width = 55
freeze_top(ws, 4)
print(f"  Audit Findings ({len(findings)}) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Anomaly Detection
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Anomaly Detection")
ws.sheet_properties.tabColor = "C00000"
write_title(ws, "Anomaly Detection Analysis", 8)

r = 3
# Analysis info
info = anomalies_data.get("analysis_info", {})
ws.cell(row=r, column=1, value="ANALYSIS SUMMARY").font = section_font
ws.merge_cells(f"A{r}:D{r}")
r += 1
for k, v in info.items():
    if isinstance(v, list):
        v = ", ".join(str(x) for x in v)
    ws.cell(row=r, column=1, value=k.replace("_", " ").title()).font = bold_font
    ws.cell(row=r, column=2, value=str(v)).font = data_font
    r += 1

# Anomaly categories
r += 1
ws.cell(row=r, column=1, value="ANOMALY CATEGORIES").font = section_font
ws.merge_cells(f"A{r}:F{r}")
r += 1
write_headers(ws, r, ["Category", "Description", "Count", "Risk Level", "Sample Transactions", ""])
r += 1
for cat in anomaly_cats:
    details = cat.get("details", [])
    sample_txns = []
    for d in details[:3]:
        refs = d.get("references", d.get("transaction_ids", d.get("affected_records", [])))
        if isinstance(refs, list):
            sample_txns.extend(refs[:2])
        elif isinstance(refs, str):
            sample_txns.append(refs)
    sample_str = ", ".join(str(s) for s in sample_txns[:5])
    write_row(ws, r, [cat.get("category", ""), cat.get("description", ""),
                       cat.get("count", ""), cat.get("risk_level", ""), sample_str])
    # Color risk level
    risk_cell = ws.cell(row=r, column=4)
    rl = cat.get("risk_level", "").lower()
    if "high" in rl or "critical" in rl:
        risk_cell.fill = critical_fill
        risk_cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    elif "medium" in rl:
        risk_cell.fill = medium_fill
    r += 1

# Suspicious entities
r += 1
ws.cell(row=r, column=1, value="SUSPICIOUS ENTITIES").font = section_font
ws.merge_cells(f"A{r}:F{r}")
r += 1
write_headers(ws, r, ["Entity ID", "Entity Name", "Type", "Risk Score", "Risk Indicators", "Recommendation"])
r += 1
for se in susp_entities:
    indicators = se.get("risk_indicators", [])
    write_row(ws, r, [se.get("entity_id", ""), se.get("entity_name", ""), se.get("entity_type", ""),
                       se.get("risk_score", ""), ", ".join(indicators) if isinstance(indicators, list) else str(indicators),
                       se.get("recommendation", "")])
    r += 1

# Fraud indicators
r += 1
ws.cell(row=r, column=1, value="FRAUD INDICATORS").font = section_font
ws.merge_cells(f"A{r}:F{r}")
r += 1
write_headers(ws, r, ["Indicator", "Category", "Risk Level", "Evidence", "Affected Area", ""])
r += 1
for fi in fraud_indicators:
    write_row(ws, r, [fi.get("indicator", ""), fi.get("category", ""), fi.get("risk_level", ""),
                       fi.get("evidence", ""), fi.get("affected_area", "")])
    r += 1

auto_width(ws)
ws.column_dimensions["B"].width = 45
ws.column_dimensions["E"].width = 50
ws.column_dimensions["F"].width = 50
freeze_top(ws, 7)
print("  Anomaly Detection ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Risk Matrix
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Risk Matrix")
ws.sheet_properties.tabColor = "843C0C"
write_title(ws, "Risk Assessment Matrix", 8)
headers = ["finding_id", "title", "severity", "risk_score", "financial_impact",
           "likelihood", "impact", "area"]
write_headers(ws, 3, headers)
# Sort by risk_score descending
sorted_findings = sorted(findings, key=lambda x: x.get("risk_score", 0), reverse=True)
for i, f in enumerate(sorted_findings):
    score = f.get("risk_score", 0)
    if score >= 16:
        likelihood = "Almost Certain"
        impact_level = "Catastrophic" if score >= 20 else "Major"
    elif score >= 12:
        likelihood = "Likely"
        impact_level = "Major"
    elif score >= 8:
        likelihood = "Possible"
        impact_level = "Moderate"
    else:
        likelihood = "Unlikely"
        impact_level = "Minor"
    r = 4 + i
    write_row(ws, r, [f.get("finding_id", ""), f.get("title", ""), f.get("severity", ""),
                       score, f.get("financial_impact", 0), likelihood, impact_level, f.get("area", "")])
    # Color by score
    score_cell = ws.cell(row=r, column=4)
    if score >= 16:
        score_cell.fill = critical_fill
        score_cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    elif score >= 12:
        score_cell.fill = high_fill
        score_cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    elif score >= 8:
        score_cell.fill = medium_fill

auto_width(ws)
ws.column_dimensions["B"].width = 50
freeze_top(ws, 4)
print(f"  Risk Matrix ({len(findings)} items) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Corrective Action Plan
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Corrective Actions")
ws.sheet_properties.tabColor = "00B050"
write_title(ws, "Corrective Action Plan", 10)

# Build corrective actions from findings
headers = ["action_id", "finding_id", "area", "severity", "action_item",
           "responsible_dept", "responsible_person", "due_date", "priority", "status"]
write_headers(ws, 3, headers)

action_id = 1
r = 4
for f in findings:
    severity = f.get("severity", "Medium")
    priority = "Immediate" if severity in ("Critical", "High") else "30 Days" if severity == "Medium" else "90 Days"
    write_row(ws, r, [
        f"CA-{action_id:03d}",
        f.get("finding_id", ""),
        f.get("area", ""),
        severity,
        f.get("recommendation", ""),
        f.get("responsible_dept", ""),
        f.get("responsible_person", ""),
        f.get("due_date", ""),
        priority,
        f.get("status", "Open")
    ])
    # Color severity
    sev_cell = ws.cell(row=r, column=4)
    if severity == "Critical":
        sev_cell.fill = critical_fill
        sev_cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    elif severity == "High":
        sev_cell.fill = high_fill
        sev_cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
    elif severity == "Medium":
        sev_cell.fill = medium_fill
    elif severity == "Low":
        sev_cell.fill = low_fill
    action_id += 1
    r += 1

auto_width(ws)
ws.column_dimensions["E"].width = 55
freeze_top(ws, 4)
print(f"  Corrective Actions ({action_id - 1} items) ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Financial Summary
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Financial Summary")
ws.sheet_properties.tabColor = "002060"
write_title(ws, "Financial Summary — Key Metrics", 6)

r = 3
# Calculate totals
total_po = sum(p.get("total_amount", 0) for p in po)
total_pi = sum(p.get("total_amount", 0) for p in pi)
total_so = sum(s.get("total_amount", 0) for s in so)
total_si = sum(s.get("total_amount", 0) for s in si)
total_expenses = sum(e.get("amount", 0) for e in expenses)
total_payroll = sum(p.get("net_salary", 0) for p in payroll)
total_tax = sum(t.get("total_paid", 0) for t in tax)

ws.cell(row=r, column=1, value="KEY FINANCIAL METRICS").font = section_font
ws.merge_cells(f"A{r}:C{r}")
r += 1
write_headers(ws, r, ["Metric", "Amount (₹)", ""])
r += 1
metrics = [
    ["Total Purchase Order Value", f"₹{total_po:,.0f}"],
    ["Total Purchase Invoice Value", f"₹{total_pi:,.0f}"],
    ["Total Sales Order Value", f"₹{total_so:,.0f}"],
    ["Total Sales Invoice Value", f"₹{total_si:,.0f}"],
    ["Total Expenses", f"₹{total_expenses:,.0f}"],
    ["Total Payroll (Net)", f"₹{total_payroll:,.0f}"],
    ["Total Tax Payments", f"₹{total_tax:,.0f}"],
    ["Total Financial Exposure from Findings", f"₹{total_impact:,.0f}"],
]
for m in metrics:
    write_row(ws, r, m)
    ws.cell(row=r, column=2).font = bold_font
    r += 1

# Budget variance
r += 1
ws.cell(row=r, column=1, value="BUDGET VARIANCE SUMMARY").font = section_font
ws.merge_cells(f"A{r}:D{r}")
r += 1
write_headers(ws, r, ["Department", "Budget (₹)", "Actual (₹)", "Variance %"])
r += 1
for b in budgets:
    write_row(ws, r, [b.get("department", ""), b.get("annual_budget", 0),
                       b.get("actual_spend", 0), b.get("variance_percent", "")])
    r += 1

auto_width(ws)
ws.column_dimensions["B"].width = 22
ws.column_dimensions["C"].width = 22
freeze_top(ws, 4)
print("  Financial Summary ✓")


# ═══════════════════════════════════════════════════════════
# SHEET: Company Profile
# ═══════════════════════════════════════════════════════════
ws = wb.create_sheet("Company Profile")
ws.sheet_properties.tabColor = "002060"
write_title(ws, "NovaTech Industries Pvt. Ltd. — Company Profile", 4)

profile = [
    ["Company Name", "NovaTech Industries Pvt. Ltd."],
    ["CIN", "U27100MH2012PTC345678"],
    ["Registered Address", "Plot No. 42, MIDC Industrial Area, Pune, Maharashtra — 411018"],
    ["Date of Incorporation", "15 March 2012"],
    ["Registered Capital", "₹10,00,00,000 (10 Crore)"],
    ["Paid-up Capital", "₹8,50,00,000 (8.5 Crore)"],
    ["GSTIN", "27AAACN1234F1Z5"],
    ["PAN", "AAACN1234F"],
    ["", ""],
    ["BUSINESS OVERVIEW", ""],
    ["Industry", "Manufacturing & Distribution"],
    ["Products", "Industrial automation components, precision instruments, electronic assemblies"],
    ["Manufacturing Plants", "3 (Pune, Chennai, Ahmedabad)"],
    ["Warehouses", "5 (Pune × 2, Chennai, Delhi, Kolkata) + 3 Plant stores"],
    ["Total Employees", "~500"],
    ["Operating States", "Maharashtra, Tamil Nadu, Gujarat, Delhi NCR, West Bengal, Karnataka, Rajasthan"],
    ["Key Customers", "BHEL, L&T, Tata Projects, Adani Power, JSW Energy, Godrej & Boyce"],
    ["Annual Revenue (FY25-26)", "~₹180 Crore"],
    ["", ""],
    ["ORGANIZATIONAL STRUCTURE", ""],
    ["Board of Directors", "Chairman: Rajesh K. Mehta, MD: Vikram S. Patel, CFO: Priya V. Iyer, COO: Amit R. Sharma"],
    ["Business Units", "BU1: Industrial Automation, BU2: Precision Instruments, BU3: Electronic Assemblies"],
    ["Key Departments", "Finance, Procurement, Production, Sales, Warehouse, Quality, HR, IT, Compliance"],
    ["", ""],
    ["AUDIT DETAILS", ""],
    ["Audit Period", "April 2025 — March 2026 (FY 2025-26)"],
    ["Audit Type", "Comprehensive Internal Audit"],
    ["Audit Team", "Lead: Suresh Nair (Internal Audit Head), Team: Kavita Desai, Rohan Gupta, Meera Joshi"],
    ["Overall Rating", "C — Needs Significant Improvement"],
    ["Total Findings", "40"],
    ["Critical Findings", "5"],
    ["Total Financial Impact", f"₹{total_impact:,.0f}"],
]

r = 3
for label, value in profile:
    if not label and not value:
        r += 1
        continue
    if not value:
        ws.cell(row=r, column=1, value=label).font = section_font
        ws.merge_cells(f"A{r}:D{r}")
        r += 1
        continue
    ws.cell(row=r, column=1, value=label).font = bold_font
    ws.cell(row=r, column=1).border = thin_border
    ws.merge_cells(f"B{r}:D{r}")
    ws.cell(row=r, column=2, value=value).font = data_font
    ws.cell(row=r, column=2).border = thin_border
    ws.cell(row=r, column=2).alignment = left_align
    r += 1

ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 25
ws.column_dimensions["C"].width = 25
ws.column_dimensions["D"].width = 25
print("  Company Profile ✓")


# ═══════════════════════════════════════════════════════════
# Final: Reorder sheets so Dashboard is first
# ═══════════════════════════════════════════════════════════
sheet_order = [
    "Dashboard", "Company Profile", "Employees", "Vendors", "Customers",
    "Products", "Warehouses", "Departments", "Purchase Orders", "Purchase Invoices",
    "Sales Orders", "Sales Invoices", "Expense Claims", "Bank Transactions",
    "Inventory Movements", "Payroll", "Tax Payments", "Budget vs Actuals",
    "Financial Summary", "User Access", "Login Activity", "Approval Workflows",
    "Audit Findings", "Anomaly Detection", "Risk Matrix", "Corrective Actions"
]
# Reorder
for i, name in enumerate(sheet_order):
    if name in wb.sheetnames:
        wb.move_sheet(name, offset=i - wb.sheetnames.index(name))

# ═══════════════════════════════════════════════════════════
# Save
# ═══════════════════════════════════════════════════════════
output = "NovaTech_Industries_Audit_FY2025-26.xlsx"
wb.save(output)
print(f"\n✅ Excel file saved: {output}")
print(f"   Sheets: {len(wb.sheetnames)}")
print(f"   Total records across all sheets: {sum([len(employees), len(vendors), len(customers), len(products), len(warehouses), len(departments), len(po), len(pi), len(so), len(si), len(expenses), len(bank), len(inv), len(payroll), len(tax), len(budgets), len(uaccess), len(login), len(findings)])}")
